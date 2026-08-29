"""МОДУЛЬ 5: Моніторинг відповідей і генерація презентації (Reply Monitor).

Періодично перевіряє Gmail-треди лідів у статусі "drafted"/"sent":
  1. Якщо чернетку користувач надіслав вручну — лід переводиться в "sent".
  2. Якщо у треді з'явилась відповідь від клієнта — генерується короткий
     план з 2-3 маркетинговими ідеями (через Claude API), на його основі
     створюється PPTX-презентація (python-pptx), і в тому ж треді
     створюється ЧЕРНЕТКА відповіді з подякою та вкладеною презентацією.
     Лід переводиться в статус "presented".

Як і решта модулів, НІКОЛИ не надсилає листи автоматично — лише читає
Gmail (users.threads.get) та створює чернетки (users.drafts.create).
"""

from __future__ import annotations

import base64
import json
import logging
import re
from dataclasses import dataclass
from datetime import UTC, datetime
from email.utils import parseaddr
from pathlib import Path
from typing import Any

from bs4 import BeautifulSoup
from pptx import Presentation

from src.clients.anthropic_client import ClaudeClient
from src.clients.gmail_client import GmailClient
from src.config import Settings
from src.db import Lead, get_leads_by_statuses, init_db, mark_sent, save_presentation, save_reply
from src.exceptions import AgentAuthError

logger = logging.getLogger(__name__)

_IDEAS_SYSTEM_PROMPT = """\
Ти — стратег performance-маркетингової агенції Vlad's Marketing. Клієнт \
відповів на наш лист холодного аутричу і зацікавився співпрацею.

На основі попереднього аналізу слабких місць бізнесу та тексту відповіді \
клієнта, запропонуй від 2 до 3 конкретних маркетингових/рекламних ідей. \
Кожна ідея має мати структуру: проблема -> рішення -> очікуваний результат.

ВАЖЛИВО:
1. Спирайся лише на надані дані (слабкі місця, текст відповіді). Не \
вигадуй цифр чи гарантій результату, яких не можна обґрунтувати.
2. "Очікуваний результат" формулюй як напрямок ефекту (наприклад, "більше \
цільових звернень з реклами"), а не як конкретний вигаданий відсоток.
3. Відповідай ВИКЛЮЧНО у форматі JSON: \
{"ideas": [{"title": "...", "problem": "...", "solution": "...", \
"expected_result": "..."}]}. Без тексту до/після JSON.
"""


@dataclass
class ReplyMonitorStats:
    """Підсумок одного запуску перевірки відповідей."""

    checked: int = 0
    marked_sent: int = 0
    presented: int = 0
    errors: int = 0

    def as_dict(self) -> dict[str, int]:
        return {
            "checked": self.checked,
            "marked_sent": self.marked_sent,
            "presented": self.presented,
            "errors": self.errors,
        }


def _b64url_decode(data: str) -> str:
    """Декодує base64url-фрагмент тіла листа Gmail у текст UTF-8."""
    padded = data + "=" * (-len(data) % 4)
    return base64.urlsafe_b64decode(padded).decode("utf-8", errors="replace")


def _extract_plain_text(payload: dict[str, Any]) -> str:
    """Рекурсивно витягує читабельний текст з payload повідомлення Gmail."""
    mime_type = payload.get("mimeType", "")
    body_data = payload.get("body", {}).get("data")

    if mime_type == "text/plain" and body_data:
        return _b64url_decode(body_data)
    if mime_type == "text/html" and body_data:
        html = _b64url_decode(body_data)
        return BeautifulSoup(html, "html.parser").get_text(separator=" ", strip=True)

    for part in payload.get("parts") or []:
        text = _extract_plain_text(part)
        if text.strip():
            return text
    return ""


def _epoch_ms_to_iso(value: str | None) -> str | None:
    if not value:
        return None
    try:
        return datetime.fromtimestamp(int(value) / 1000, tz=UTC).isoformat()
    except (ValueError, TypeError):
        return None


def _classify_thread(messages: list[dict[str, Any]], sender_email: str) -> dict[str, Any]:
    """Аналізує повідомлення треду: чи лист надіслано, чи є відповідь клієнта.

    Args:
        messages: список повідомлень з `users.threads.get(format="full")`.
        sender_email: власна email-адреса користувача (для розрізнення
            "наш надісланий лист" від "відповідь клієнта").

    Returns:
        Словник з ключами: sent (bool), sent_at (str|None),
        reply_text (str|None), reply_message_id (str|None), reply_from (str|None).
    """
    sender_lower = (sender_email or "").strip().lower()
    sent = False
    sent_at: str | None = None
    reply_text: str | None = None
    reply_message_id: str | None = None
    reply_from: str | None = None
    latest_reply_ts = -1

    for msg in messages:
        headers = {
            h["name"].lower(): h["value"]
            for h in msg.get("payload", {}).get("headers", [])
            if "name" in h
        }
        from_email = parseaddr(headers.get("from", ""))[1].lower()
        label_ids = msg.get("labelIds", []) or []
        internal_date = msg.get("internalDate")

        is_from_us = bool(sender_lower) and from_email == sender_lower
        if is_from_us or "SENT" in label_ids:
            sent = True
            if sent_at is None:
                sent_at = _epoch_ms_to_iso(internal_date)
            continue

        # Повідомлення не від нас -> це відповідь клієнта. Беремо найновішу.
        text = _extract_plain_text(msg.get("payload", {}))
        if not text.strip():
            continue
        try:
            ts = int(internal_date) if internal_date else 0
        except ValueError:
            ts = 0
        if ts >= latest_reply_ts:
            latest_reply_ts = ts
            reply_text = text.strip()
            reply_message_id = headers.get("message-id")
            reply_from = parseaddr(headers.get("from", ""))[1] or headers.get("from")

    return {
        "sent": sent,
        "sent_at": sent_at,
        "reply_text": reply_text,
        "reply_message_id": reply_message_id,
        "reply_from": reply_from,
    }


def parse_ideas_response(response_text: str) -> list[dict[str, str]]:
    """Розпарсює відповідь Claude у список ідей {title, problem, solution, expected_result}."""
    text = response_text.strip()
    try:
        data = json.loads(text)
    except json.JSONDecodeError as exc:
        match = re.search(r"\{.*\}", text, re.DOTALL)
        if not match:
            raise ValueError(
                f"Не вдалося розпарсити ідеї з відповіді Claude: {text[:200]!r}"
            ) from exc
        data = json.loads(match.group(0))

    ideas = data.get("ideas") if isinstance(data, dict) else None
    if not isinstance(ideas, list) or not ideas:
        raise ValueError("Відповідь Claude не містить непорожнього списку 'ideas'.")

    cleaned = []
    for idea in ideas[:3]:
        if not isinstance(idea, dict):
            continue
        cleaned.append(
            {
                "title": str(idea.get("title", "Ідея")).strip(),
                "problem": str(idea.get("problem", "")).strip(),
                "solution": str(idea.get("solution", "")).strip(),
                "expected_result": str(idea.get("expected_result", "")).strip(),
            }
        )
    if not cleaned:
        raise ValueError("Жодної коректної ідеї не знайдено у відповіді Claude.")
    return cleaned


def generate_ideas(claude: ClaudeClient, lead: Lead, reply_text: str) -> list[dict[str, str]]:
    """Викликає Claude API, щоб згенерувати 2-3 маркетингові ідеї для ліда.

    Використовує дані аналізу з Модуля 3 (weaknesses) та текст відповіді
    клієнта як контекст.
    """
    weaknesses_text = "\n".join(f"- {w}" for w in lead.weaknesses) or "(немає даних аналізу)"
    prompt = (
        f"Бізнес: {lead.name} ({lead.category or 'категорія невідома'})\n\n"
        f"Слабкі місця з попереднього аналізу:\n{weaknesses_text}\n\n"
        f'Текст відповіді клієнта на наш лист:\n"""\n{reply_text}\n"""\n\n'
        "Запропонуй 2-3 маркетингові ідеї для цього бізнесу."
    )
    response_text = claude.complete(
        system=_IDEAS_SYSTEM_PROMPT,
        user_prompt=prompt,
        max_tokens=900,
        dry_run_label=f"Генерація ідей для '{lead.name}'",
    )
    return parse_ideas_response(response_text)


def generate_presentation(lead: Lead, ideas: list[dict[str, str]], presentations_dir: str) -> str:
    """Генерує PPTX-презентацію на основі маркетингових ідей.

    Структура: title slide -> по одному слайду на ідею (проблема -> рішення
    -> очікуваний результат) -> фінальний слайд з CTA.

    Args:
        lead: лід, для якого генерується презентація.
        ideas: список ідей з generate_ideas()/parse_ideas_response().
        presentations_dir: директорія для збереження .pptx файлів.

    Returns:
        Шлях до збереженого .pptx файлу.
    """
    prs = Presentation()

    title_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_layout)
    slide.shapes.title.text = lead.name
    slide.placeholders[1].text = "Пропозиція з ідеями для просування — Vlad's Marketing"

    content_layout = prs.slide_layouts[1]
    for idea in ideas:
        slide = prs.slides.add_slide(content_layout)
        slide.shapes.title.text = idea["title"]
        body = slide.placeholders[1].text_frame
        body.text = f"Проблема: {idea['problem']}"
        p_solution = body.add_paragraph()
        p_solution.text = f"Рішення: {idea['solution']}"
        p_result = body.add_paragraph()
        p_result.text = f"Очікуваний результат: {idea['expected_result']}"

    cta_slide = prs.slides.add_slide(content_layout)
    cta_slide.shapes.title.text = "Готові розпочати?"
    cta_body = cta_slide.placeholders[1].text_frame
    cta_body.text = "Обговорімо деталі та узгодимо перші кроки кампанії."
    p_contact = cta_body.add_paragraph()
    p_contact.text = "З повагою, Владислав, Vlad's Marketing"

    out_dir = Path(presentations_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    safe_name = re.sub(r"[^\w\-]+", "_", lead.name).strip("_") or "lead"
    out_path = out_dir / f"presentation_{lead.id}_{safe_name}_{timestamp}.pptx"
    prs.save(str(out_path))
    return str(out_path)


def build_reply_draft_body(lead: Lead, ideas: list[dict[str, str]]) -> str:
    """Формує текст чернетки-подяки з коротким описом презентації у вкладенні."""
    idea_titles = "\n".join(f"- {idea['title']}" for idea in ideas)
    return (
        f"Доброго дня!\n\n"
        f"Дякую за відповідь! Підготував короткий план із {len(ideas)} ідеями "
        f"для {lead.name} — деталі у презентації у вкладенні:\n\n"
        f"{idea_titles}\n\n"
        "Готовий обговорити деталі та наступні кроки у зручний для вас час.\n\n"
        "З повагою,\nВладислав, Vlad's Marketing"
    )


def check_replies(settings: Settings, *, dry_run: bool = False) -> ReplyMonitorStats:
    """Перевіряє Gmail-треди лідів у статусах "drafted"/"sent" на нові події.

    Args:
        settings: налаштування агента.
        dry_run: якщо True — не звертається реально до Gmail/Claude API.

    Returns:
        ReplyMonitorStats зі статистикою запуску.
    """
    init_db(settings.leads_db_path)
    stats = ReplyMonitorStats()

    leads = get_leads_by_statuses(settings.leads_db_path, ("drafted", "sent"))
    if not leads:
        logger.warning(
            "Немає лідів у статусі 'drafted' або 'sent' для перевірки відповідей. "
            "Спочатку створіть чернетки через 'python main.py draft' і надішліть їх вручну."
        )
        return stats

    gmail = GmailClient(settings.gmail_credentials_file, settings.gmail_token_file, dry_run=dry_run)
    claude = ClaudeClient(settings.anthropic_api_key, settings.anthropic_model, dry_run=dry_run)

    logger.info("Старт Модуля 5 (Reply Monitor): перевірка %d тредів.", len(leads))

    for lead in leads:
        stats.checked += 1
        if not lead.thread_id:
            logger.warning("Лід #%d ('%s') не має thread_id — пропущено.", lead.id, lead.name)
            continue

        if dry_run:
            logger.info("[DRY-RUN] Було б перевірено тред ліда #%d ('%s').", lead.id, lead.name)
            print(f"[DRY-RUN] Reply Monitor -> тред '{lead.name}' не перевіряється реально.")
            continue

        try:
            thread = gmail.get_thread(lead.thread_id)
        except AgentAuthError:
            raise
        except Exception as exc:  # noqa: BLE001
            stats.errors += 1
            logger.error("Не вдалося отримати тред ліда #%d ('%s'): %s", lead.id, lead.name, exc)
            continue

        classification = _classify_thread(thread.get("messages", []), settings.gmail_sender_email)

        current_status = lead.status
        if current_status == "drafted" and classification["sent"]:
            mark_sent(settings.leads_db_path, lead.id, classification["sent_at"])
            stats.marked_sent += 1
            current_status = "sent"
            logger.info(
                "Лід #%d ('%s') позначено як 'sent' (лист надіслано вручну).", lead.id, lead.name
            )

        if current_status == "sent" and classification["reply_text"]:
            logger.info(
                "Виявлено відповідь від ліда #%d ('%s'). Генеруємо презентацію.", lead.id, lead.name
            )
            try:
                ideas = generate_ideas(claude, lead, classification["reply_text"])
                pptx_path = generate_presentation(lead, ideas, settings.presentations_dir)
                reply_body = build_reply_draft_body(lead, ideas)
                recipient = classification["reply_from"] or lead.enrichment.get("contact_email")
                draft = gmail.create_draft(
                    to=recipient,
                    subject=f"Re: {lead.email_subject or lead.name}",
                    body_text=reply_body,
                    thread_id=lead.thread_id,
                    attachment_path=pptx_path,
                    in_reply_to_message_id=classification["reply_message_id"],
                )
            except AgentAuthError:
                raise
            except Exception as exc:  # noqa: BLE001
                stats.errors += 1
                logger.error(
                    "Не вдалося згенерувати презентацію/чернетку відповіді для ліда #%d ('%s'): %s",
                    lead.id,
                    lead.name,
                    exc,
                )
                continue

            save_reply(settings.leads_db_path, lead.id, classification["reply_text"])
            save_presentation(
                settings.leads_db_path,
                lead.id,
                presentation_path=pptx_path,
                reply_draft_id=draft["id"],
            )
            stats.presented += 1
            logger.info(
                "Лід #%d ('%s') переведено у статус 'presented'. Презентація: %s",
                lead.id,
                lead.name,
                pptx_path,
            )

    logger.info(
        "Модуль 5 завершено: %d перевірено, %d позначено 'sent', %d представлено, %d помилок.",
        stats.checked,
        stats.marked_sent,
        stats.presented,
        stats.errors,
    )
    return stats
