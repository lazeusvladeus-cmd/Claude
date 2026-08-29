"""Тести Модуля 5 (reply_monitor.py): моніторинг відповідей + генерація PPTX."""

from __future__ import annotations

import base64
import json

import pytest
from pptx import Presentation

import reply_monitor
from src.db import get_lead, save_analysis, save_draft, save_enrichment
from src.exceptions import AgentAuthError
from tests.conftest import load_fixture


def _b64(text: str) -> str:
    return base64.urlsafe_b64encode(text.encode("utf-8")).decode("ascii")


def _make_message(
    *, from_addr: str, text: str, label_sent: bool, internal_date_ms: int, message_id: str
) -> dict:
    return {
        "id": message_id,
        "labelIds": ["SENT"] if label_sent else ["INBOX"],
        "internalDate": str(internal_date_ms),
        "payload": {
            "mimeType": "text/plain",
            "headers": [
                {"name": "From", "value": from_addr},
                {"name": "Message-ID", "value": f"<{message_id}@example.ua>"},
            ],
            "body": {"data": _b64(text)},
        },
    }


def _drafted_lead(sample_lead_factory, settings, *, place_id: str, name: str, email: str):
    lead_id = sample_lead_factory(
        place_id=place_id, name=name, website="https://example.ua", social_links="[]"
    )
    save_enrichment(settings.leads_db_path, lead_id, {"source": "website", "contact_email": email})
    save_analysis(settings.leads_db_path, lead_id, ["Немає блогу.", "Немає онлайн-замовлення."])
    save_draft(
        settings.leads_db_path,
        lead_id,
        draft_id="draft-1",
        thread_id="thread-1",
        subject="Пропозиція співпраці",
        body="Текст листа",
    )
    return lead_id


class TestClassifyThread:
    def test_detects_sent_message(self) -> None:
        messages = [
            _make_message(
                from_addr="vlad@vladsmarketing.example",
                text="Наш лист",
                label_sent=True,
                internal_date_ms=1000,
                message_id="m1",
            )
        ]
        result = reply_monitor._classify_thread(messages, "vlad@vladsmarketing.example")
        assert result["sent"] is True
        assert result["reply_text"] is None

    def test_detects_reply_from_client(self) -> None:
        messages = [
            _make_message(
                from_addr="vlad@vladsmarketing.example",
                text="Наш лист",
                label_sent=True,
                internal_date_ms=1000,
                message_id="m1",
            ),
            _make_message(
                from_addr="client@business.example",
                text="Дякую, цікаво! Розкажіть детальніше.",
                label_sent=False,
                internal_date_ms=2000,
                message_id="m2",
            ),
        ]
        result = reply_monitor._classify_thread(messages, "vlad@vladsmarketing.example")
        assert result["sent"] is True
        assert result["reply_text"] == "Дякую, цікаво! Розкажіть детальніше."
        assert result["reply_from"] == "client@business.example"
        assert result["reply_message_id"] == "<m2@example.ua>"

    def test_picks_latest_reply_when_multiple(self) -> None:
        messages = [
            _make_message(
                from_addr="client@business.example",
                text="Перше",
                label_sent=False,
                internal_date_ms=2000,
                message_id="m2",
            ),
            _make_message(
                from_addr="client@business.example",
                text="Друге, новіше",
                label_sent=False,
                internal_date_ms=3000,
                message_id="m3",
            ),
        ]
        result = reply_monitor._classify_thread(messages, "vlad@vladsmarketing.example")
        assert result["reply_text"] == "Друге, новіше"

    def test_no_reply_yet(self) -> None:
        messages = [
            _make_message(
                from_addr="vlad@vladsmarketing.example",
                text="Наш лист",
                label_sent=True,
                internal_date_ms=1000,
                message_id="m1",
            )
        ]
        result = reply_monitor._classify_thread(messages, "vlad@vladsmarketing.example")
        assert result["reply_text"] is None


class TestParseIdeasResponse:
    def test_parses_clean_json(self) -> None:
        fixture = load_fixture("claude_ideas_response.json")
        ideas = reply_monitor.parse_ideas_response(json.dumps(fixture, ensure_ascii=False))
        assert len(ideas) == 2
        assert ideas[0]["title"] == fixture["ideas"][0]["title"]

    def test_raises_on_missing_ideas_key(self) -> None:
        with pytest.raises(ValueError):
            reply_monitor.parse_ideas_response(json.dumps({"foo": "bar"}))

    def test_limits_to_three_ideas(self) -> None:
        data = {
            "ideas": [
                {"title": f"Ідея {i}", "problem": "п", "solution": "р", "expected_result": "о"}
                for i in range(5)
            ]
        }
        ideas = reply_monitor.parse_ideas_response(json.dumps(data))
        assert len(ideas) == 3


class TestGeneratePresentation:
    def test_pptx_has_correct_structure(self, settings, sample_lead_factory) -> None:
        lead_id = sample_lead_factory(place_id="p1", name="Кав'ярня Комаха")
        lead = get_lead(settings.leads_db_path, lead_id)
        ideas = load_fixture("claude_ideas_response.json")["ideas"]

        path = reply_monitor.generate_presentation(lead, ideas, settings.presentations_dir)

        prs = Presentation(path)
        # Title slide + одна на кожну ідею + CTA slide.
        assert len(prs.slides) == len(ideas) + 2
        title_slide = prs.slides[0]
        assert lead.name in title_slide.shapes.title.text

        first_idea_slide = prs.slides[1]
        assert ideas[0]["title"] == first_idea_slide.shapes.title.text
        body_text = first_idea_slide.placeholders[1].text_frame.text
        assert "Проблема" in body_text
        assert "Рішення" in body_text
        assert "Очікуваний результат" in body_text

        cta_slide = prs.slides[-1]
        assert "Готові розпочати" in cta_slide.shapes.title.text


class TestCheckReplies:
    def test_marks_sent_when_draft_was_actually_sent(
        self, settings, sample_lead_factory, mocker
    ) -> None:
        lead_id = _drafted_lead(
            sample_lead_factory, settings, place_id="p2", name="Комаха", email="hello@example.ua"
        )
        thread = {
            "messages": [
                _make_message(
                    from_addr=settings.gmail_sender_email,
                    text="Наш лист",
                    label_sent=True,
                    internal_date_ms=1000,
                    message_id="m1",
                )
            ]
        }
        mocker.patch("reply_monitor.GmailClient.get_thread", return_value=thread)

        stats = reply_monitor.check_replies(settings, dry_run=False)

        assert stats.marked_sent == 1
        assert stats.presented == 0
        assert get_lead(settings.leads_db_path, lead_id).status == "sent"

    def test_generates_presentation_and_reply_draft_on_reply(
        self, settings, sample_lead_factory, mocker
    ) -> None:
        lead_id = _drafted_lead(
            sample_lead_factory, settings, place_id="p3", name="Комаха", email="hello@example.ua"
        )
        from src.db import mark_sent

        mark_sent(settings.leads_db_path, lead_id)

        thread = {
            "messages": [
                _make_message(
                    from_addr=settings.gmail_sender_email,
                    text="Наш лист",
                    label_sent=True,
                    internal_date_ms=1000,
                    message_id="m1",
                ),
                _make_message(
                    from_addr="hello@example.ua",
                    text="Дякую, цікаво! Надішліть план.",
                    label_sent=False,
                    internal_date_ms=2000,
                    message_id="m2",
                ),
            ]
        }
        mocker.patch("reply_monitor.GmailClient.get_thread", return_value=thread)
        ideas_fixture = load_fixture("claude_ideas_response.json")
        mocker.patch(
            "reply_monitor.ClaudeClient.complete",
            return_value=json.dumps(ideas_fixture, ensure_ascii=False),
        )

        mock_service = mocker.MagicMock()
        drafts_create = mock_service.users.return_value.drafts.return_value.create
        drafts_create.return_value.execute.return_value = {
            "id": "reply-draft-1",
            "message": {"threadId": "thread-1"},
        }
        mocker.patch("reply_monitor.GmailClient._get_service", return_value=mock_service)

        stats = reply_monitor.check_replies(settings, dry_run=False)

        assert stats.presented == 1
        lead = get_lead(settings.leads_db_path, lead_id)
        assert lead.status == "presented"
        assert lead.reply_text == "Дякую, цікаво! Надішліть план."
        assert lead.presentation_path is not None

        from pathlib import Path

        assert Path(lead.presentation_path).exists()

        # send() ніколи не викликається, лише drafts().create().
        mock_service.users.return_value.messages.return_value.send.assert_not_called()
        mock_service.users.return_value.drafts.return_value.create.assert_called_once()

    def test_empty_db_is_graceful(self, settings) -> None:
        stats = reply_monitor.check_replies(settings, dry_run=False)
        assert stats.checked == 0

    def test_auth_error_propagates(self, settings, sample_lead_factory, mocker) -> None:
        _drafted_lead(sample_lead_factory, settings, place_id="p4", name="X", email="x@example.ua")
        mocker.patch(
            "reply_monitor.GmailClient.get_thread",
            side_effect=AgentAuthError("Токен протерміновано."),
        )

        with pytest.raises(AgentAuthError):
            reply_monitor.check_replies(settings, dry_run=False)

    def test_dry_run_makes_no_real_calls(self, settings, sample_lead_factory, mocker) -> None:
        _drafted_lead(sample_lead_factory, settings, place_id="p5", name="Y", email="y@example.ua")
        gmail_spy = mocker.patch("reply_monitor.GmailClient.get_thread")
        claude_spy = mocker.patch("reply_monitor.ClaudeClient.complete")

        stats = reply_monitor.check_replies(settings, dry_run=True)

        gmail_spy.assert_not_called()
        claude_spy.assert_not_called()
        assert stats.presented == 0
