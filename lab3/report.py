"""Збереження результатів у output/: таблиці .csv, графік .png і зведений .xlsx."""

from __future__ import annotations

from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from task_3_2_market import plot_market

FMT_MONEY = "#,##0.00"
FMT_SOLUTION = "0.0000"
FMT_INT = "0"

_BOLD = Font(bold=True)
_TITLE = Font(bold=True, size=12)
_HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")
_THIN = Side(style="thin", color="999999")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def save_csv(tables: dict[str, pd.DataFrame], out_dir: Path) -> list[Path]:
    """Зберігає таблиці у CSV (UTF-8 з BOM, щоб Excel коректно показав кирилицю)."""
    paths = []
    for name, df in tables.items():
        path = out_dir / f"{name}.csv"
        df.to_csv(path, index=False, encoding="utf-8-sig")
        paths.append(path)
    return paths


def _write_table(ws, row: int, title: str, df: pd.DataFrame,
                 formats: dict[str, str] | None = None, col: int = 1) -> int:
    """Пише таблицю із заголовком на аркуш і повертає номер наступного вільного рядка."""
    formats = formats or {}
    ws.cell(row=row, column=col, value=title).font = _TITLE
    row += 1
    for j, name in enumerate(df.columns):
        c = ws.cell(row=row, column=col + j, value=str(name))
        c.font, c.fill, c.border = _BOLD, _HEADER_FILL, _BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for values in df.itertuples(index=False):
        row += 1
        for j, (name, value) in enumerate(zip(df.columns, values)):
            if isinstance(value, np.generic):
                value = value.item()
            c = ws.cell(row=row, column=col + j, value=value)
            c.border = _BORDER
            if name in formats:
                c.number_format = formats[name]
    return row + 2


def _write_matrix(ws, row: int, title: str, M: np.ndarray, fmt: str,
                  labels: list[str] | None = None) -> int:
    """Пише матрицю/вектор з підписами стовпців."""
    M = np.atleast_2d(M)
    labels = labels or [f"Стовпець {j + 1}" for j in range(M.shape[1])]
    df = pd.DataFrame(M, columns=labels)
    return _write_table(ws, row, title, df, {c: fmt for c in labels})


def _set_widths(ws, widths: dict[int, float]) -> None:
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width


def build_workbook(r31: dict, r32: dict, r33: dict, k: int, surname: str,
                   path: Path) -> None:
    """Формує зведений файл Excel: Аркуш1 — 3.1, Аркуш2 — 3.2, Аркуш3 — 3.3."""
    wb = Workbook()
    money = {"Вартість поставленого товару": FMT_MONEY, "Ціна": FMT_MONEY,
             "Кількість поставленого товару": FMT_INT}

    # Аркуш1 — завдання 3.1
    ws = wb.active
    ws.title = "Аркуш1"
    ws.cell(row=1, column=1, value=f"Лабораторна робота №3. Завдання 3.1. Виконав(ла): {surname}, k = {k}").font = Font(bold=True, size=14)
    row = 3
    row = _write_table(ws, row, "Таблиця 1. Постачання товарів", r31["table_1"], money)
    row = _write_table(ws, row, "Ціни за одиницю товару (аналог Sheet2)", r31["prices"], money)
    row = _write_table(ws, row, "Таблиця 7. Товар Т1 у магазин М1 (умова І)", r31["table_7"], money)
    row = _write_table(ws, row, "Таблиця 8. Товари Т1 або Т4 (умова АБО)", r31["table_8"], money)
    _write_table(ws, row, "Таблиця 9. Магазини без повторень", r31["table_9"])
    _set_widths(ws, {1: 12, 2: 12, 3: 12, 4: 18, 5: 18})
    for col in (4, 5):
        ws.cell(row=4, column=col).alignment = Alignment(wrap_text=True, horizontal="center")

    # Аркуш2 — завдання 3.2
    ws = wb.create_sheet("Аркуш2")
    ws.cell(row=1, column=1, value=f"Завдання 3.2. Попит, пропозиція, рівновага. {surname}, k = {k}").font = Font(bold=True, size=14)
    row = _write_table(ws, 3, "Таблиця 10. Попит і пропозиція", r32["table_10"],
                       {"Кількість товару": FMT_INT, "Попит": FMT_MONEY, "Пропозиція": FMT_MONEY})
    start = row
    row = _write_table(ws, row, "Таблиця 11. Підбір параметра (рівновага)", r32["table_11"])
    # Q* — 4 знаки, ціни — 2 знаки, різниця — експоненційний формат.
    for i, fmt in enumerate([FMT_SOLUTION, FMT_MONEY, FMT_MONEY, "0.00E+00", FMT_MONEY]):
        ws.cell(row=start + 2 + i, column=2).number_format = fmt
    ws.cell(row=row, column=1, value="Висновок").font = _TITLE
    c = ws.cell(row=row + 1, column=1, value=r32["conclusion"])
    c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=row + 1, start_column=1, end_row=row + 6, end_column=4)
    _set_widths(ws, {1: 32, 2: 16, 3: 16, 4: 16})

    # Аркуш3 — завдання 3.3
    ws = wb.create_sheet("Аркуш3")
    ws.cell(row=1, column=1, value=f"Завдання 3.3. Система лінійних рівнянь. {surname}, k = {k}").font = Font(bold=True, size=14)
    xcols = ["x1", "x2", "x3"]
    row = _write_matrix(ws, 3, "Матриця A", r33["A"], FMT_INT, xcols)
    row = _write_matrix(ws, row, "Вектор b", r33["b"].reshape(-1, 1), FMT_INT, ["b"])
    row = _write_matrix(ws, row, "Обернена матриця A⁻¹ (MINVERSE)", r33["A_inv"], FMT_SOLUTION, xcols)
    row = _write_matrix(ws, row, "Розв'язок x = A⁻¹·b (MMULT)", r33["x_inverse"], FMT_SOLUTION, xcols)
    for i, Ai in enumerate(r33["cramer_matrices"], start=1):
        row = _write_matrix(ws, row, f"Допоміжна матриця A{i}", Ai, FMT_INT, xcols)
    dets = pd.DataFrame({
        "Визначник": ["Δ", "Δ1", "Δ2", "Δ3"],
        "Значення (MDETERM)": [r33["delta"], *r33["deltas"]],
    })
    row = _write_table(ws, row, "Визначники", dets, {"Значення (MDETERM)": FMT_SOLUTION})
    row = _write_matrix(ws, row, "Розв'язок методом Крамера xᵢ = Δᵢ/Δ", r33["x_cramer"], FMT_SOLUTION, xcols)
    check = pd.DataFrame({
        "b": r33["b"],
        "A·x (обернена)": r33["A"] @ r33["x_inverse"],
        "A·x (Крамер)": r33["A"] @ r33["x_cramer"],
    })
    _write_table(ws, row, "Перевірка A·x = b", check, {c: FMT_SOLUTION for c in check.columns})
    _set_widths(ws, {1: 20, 2: 20, 3: 20})

    wb.save(path)


def save_all(r31: dict, r32: dict, r33: dict, k: int, surname: str,
             out_dir: Path) -> dict[str, Path]:
    """Зберігає всі результати і повертає шляхи до створених файлів."""
    out_dir.mkdir(parents=True, exist_ok=True)
    tables = {
        "table_1": r31["table_1"],
        "table_7": r31["table_7"],
        "table_8": r31["table_8"],
        "table_9": r31["table_9"],
        "table_10": r32["table_10"],
        "table_11": r32["table_11"],
    }
    paths = {p.stem: p for p in save_csv(tables, out_dir)}
    paths["market"] = out_dir / "market.png"
    plot_market(r32["table_10"], r32["equilibrium"], surname, paths["market"])
    paths["xlsx"] = out_dir / f"{surname}_Лаб3.xlsx"
    build_workbook(r31, r32, r33, k, surname, paths["xlsx"])
    return paths
