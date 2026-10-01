"""Тести лабораторної №3: очікувані результати для k = 7 і незалежна перевірка для інших k."""

import math

import numpy as np
import pytest
from openpyxl import load_workbook
from scipy.optimize import brentq

import main
import task_3_1_filter as t31
import task_3_2_market as t32
import task_3_3_linear as t33

K = 7
OTHER_K = [1, 3, 11, 25]


# ---------- Завдання 3.1 ----------

def test_table_1_costs_k7():
    df = t31.build_table_1(K)
    assert df[t31.COST].tolist() == [170, 1200, 85, 255, 210, 1140, 405, 680]
    assert df.loc[4, t31.QTY] == 14
    assert df.loc[5, t31.QTY] == 57


def test_table_7_k7():
    df = t31.table_7(t31.build_table_1(K))
    rows = [tuple(r) for r in df.iloc[:, :4].itertuples(index=False)]
    assert rows == [("Б1", "М1", "Т1", 10), ("Б3", "М1", "Т1", 15), ("Б4", "М1", "Т1", 40)]


def test_table_8_k7():
    df = t31.table_8(t31.build_table_1(K))
    rows = [tuple(r) for r in df[["База", "Магазин", "Товар"]].itertuples(index=False)]
    assert rows == [
        ("Б1", "М1", "Т1"), ("Б1", "М1", "Т4"), ("Б1", "М2", "Т1"),
        ("Б3", "М1", "Т1"), ("Б3", "М2", "Т4"), ("Б4", "М1", "Т1"),
    ]


def test_table_9_k7():
    assert t31.table_9(t31.build_table_1(K))["Магазин"].tolist() == ["М1", "М2"]


@pytest.mark.parametrize("k", OTHER_K)
def test_table_1_other_k(k):
    df = t31.build_table_1(k)
    prices = {"Т1": k + 10, "Т2": 15, "Т4": 20, "Т5": k + 20}
    for _, row in df.iterrows():
        assert row[t31.COST] == row[t31.QTY] * prices[row["Товар"]]
    assert df.loc[4, t31.QTY] == 2 * k
    assert df.loc[5, t31.QTY] == k + 50


@pytest.mark.parametrize("k", [K, *OTHER_K])
def test_filters_match_query(k):
    df = t31.build_table_1(k)
    q = t31.query_examples(df)
    assert t31.table_7(df).equals(q['Магазин == "М1" and Товар == "Т1"'])
    assert t31.table_8(df).equals(q['Товар in ["Т1", "Т4"]'])


# ---------- Завдання 3.2 ----------

def test_table_10_k7():
    df = t32.build_table_10(K).set_index("Кількість товару")
    assert len(df) == 15
    assert df.loc[10, "Попит"] == pytest.approx(700.00, abs=5e-3)
    assert df.loc[50, "Попит"] == pytest.approx(140.00, abs=5e-3)
    assert df.loc[80, "Попит"] == pytest.approx(87.50, abs=5e-3)
    assert df.loc[10, "Пропозиція"] == pytest.approx(2.72, abs=5e-3)
    assert df.loc[50, "Пропозиція"] == pytest.approx(148.41, abs=5e-3)
    assert df.loc[80, "Пропозиція"] == pytest.approx(2980.96, abs=5e-3)


def test_equilibrium_k7():
    eq = t32.find_equilibrium(K)
    assert eq.q == pytest.approx(49.5141, abs=1e-3)
    assert eq.price == pytest.approx(141.3739, abs=1e-2)
    assert abs(eq.diff) < 1e-6
    assert t32.crossing_interval(t32.build_table_10(K)) == (45, 50)


@pytest.mark.parametrize("k", OTHER_K)
def test_equilibrium_other_k(k):
    expected = brentq(lambda q: 1000 * k / q - math.exp(q / 10), 10, 80)
    r = t32.run(k)
    eq = r["equilibrium"]
    assert eq.q == pytest.approx(expected, abs=1e-6)
    assert eq.price == pytest.approx(math.exp(expected / 10), abs=1e-6)
    assert abs(eq.diff) < 1e-6
    assert r["q_fsolve"] == pytest.approx(expected, abs=1e-6)
    assert r["q_bisection"] == pytest.approx(expected, abs=1e-6)
    lo, hi = r["interval"]
    assert hi - lo == 5 and lo <= expected <= hi


# ---------- Завдання 3.3 ----------

def test_linear_k7():
    r = t33.run(K)
    assert r["delta"] == pytest.approx(102)
    assert r["deltas"] == pytest.approx([822, 924, 1158])
    expected = [8.0588, 9.0588, 11.3529]
    assert r["x_inverse"] == pytest.approx(expected, abs=1e-4)
    assert r["x_cramer"] == pytest.approx(expected, abs=1e-4)
    assert r["A_inv"][0] == pytest.approx([0.1765, 0.0, 0.0588], abs=1e-4)
    assert r["check_inverse"] and r["check_cramer"] and r["methods_agree"]


@pytest.mark.parametrize("k", OTHER_K)
def test_linear_other_k(k):
    A = np.array(t33.DEFAULT_A, dtype=float)
    b = np.array([38, 6 * k + 25, 30 - k], dtype=float)
    expected = np.linalg.solve(A, b)
    r = t33.run(k)
    np.testing.assert_allclose(r["x_inverse"], expected, atol=1e-9)
    np.testing.assert_allclose(r["x_cramer"], expected, atol=1e-9)
    np.testing.assert_allclose(A @ r["x_cramer"], b, atol=1e-9)


def test_cramer_matrices_replace_column():
    A, b = t33.build_system(K)
    for i, Ai in enumerate(t33.cramer_matrices(A, b)):
        np.testing.assert_array_equal(Ai[:, i], b)
        np.testing.assert_array_equal(np.delete(Ai, i, axis=1), np.delete(A, i, axis=1))


def test_manual_system_replaceable():
    r = t33.run(K, A=t33.MANUAL_A, b=t33.manual_b(K))
    np.testing.assert_allclose(r["A"] @ r["x_cramer"], t33.manual_b(K), atol=1e-9)


def test_singular_matrix_rejected():
    with pytest.raises(np.linalg.LinAlgError):
        t33.solve_inverse(np.array([[1.0, 2.0], [2.0, 4.0]]), np.array([1.0, 2.0]))


# ---------- Звіт ----------

@pytest.mark.parametrize("k", [K, 11])
def test_main_creates_outputs(tmp_path, k, capsys):
    main.main(["--k", str(k), "--surname", "Колеснік", "--output", str(tmp_path)])
    names = [f"table_{n}.csv" for n in (1, 7, 8, 9, 10, 11)] + ["market.png", "Колеснік_Лаб3.xlsx"]
    for name in names:
        assert (tmp_path / name).stat().st_size > 0, name
    wb = load_workbook(tmp_path / "Колеснік_Лаб3.xlsx")
    assert wb.sheetnames == ["Аркуш1", "Аркуш2", "Аркуш3"]
    out = capsys.readouterr().out
    for section in ("ЗАВДАННЯ 3.1", "ЗАВДАННЯ 3.2", "ЗАВДАННЯ 3.3"):
        assert section in out
    assert f"k = {k}" in out
